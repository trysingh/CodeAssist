import asyncio
import base64
import io
import re
import sys
from typing import Any, Dict, List, Tuple

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_ollama import ChatOllama
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config.settings import settings

MAX_TEXT_CHARS = 300_000      # hard cap on text kept per attachment
OPENING_CHARS = 1200          # opening excerpt shown for large attachments

OCR_PROMPT = (
    "Transcribe ALL text visible in this image exactly as written, preserving line breaks, "
    "code indentation and table structure (use markdown tables). "
    "If the image is a diagram, UI screenshot or chart, then after the transcription add a short "
    "'Description:' section covering components, labels, relationships and flow. "
    "Output only the extracted content, no commentary."
)


class AttachmentService:
    """Turns user attachments (image / pdf / docx / text) into text and serves it back per chat thread."""

    def __init__(self, vector_service):
        self.embeddings = vector_service.embeddings
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=150)

    # ---------------- extraction ----------------
    @staticmethod
    def classify(filename: str, mime: str) -> str:
        name, mime = (filename or "").lower(), (mime or "").lower()
        if mime.startswith("image/") or name.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp")):
            return "image"
        if mime == "application/pdf" or name.endswith(".pdf"):
            return "pdf"
        if name.endswith(".docx"):
            return "docx"
        return "text"

    async def extract(self, filename: str, mime: str, data: bytes) -> Tuple[str, str]:
        kind = self.classify(filename, mime)
        if kind == "image":
            text = await self._ocr_image(data, mime or "image/png")
        elif kind == "pdf":
            text = await asyncio.to_thread(self._read_pdf, data)
        elif kind == "docx":
            text = await asyncio.to_thread(self._read_docx, data)
        else:
            text = self._read_text(filename, data)
        return kind, text.replace("\x00", "").strip()[:MAX_TEXT_CHARS]

    async def _ocr_image(self, data: bytes, mime: str) -> str:
        llm = ChatOllama(
            model=settings.vision_model or settings.llm_model,
            base_url=settings.ollama_base_url,
            temperature=0,
        )
        b64 = base64.b64encode(data).decode()
        msg = HumanMessage(content=[
            {"type": "text", "text": OCR_PROMPT},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
        ])
        res = await llm.ainvoke([msg])
        return str(res.content)

    @staticmethod
    def _read_pdf(data: bytes) -> str:
        from pypdf import PdfReader
        pages = [(p.extract_text() or "").strip() for p in PdfReader(io.BytesIO(data)).pages]
        if not any(pages):
            return "[PDF has no extractable text layer (likely a scan).]"
        return "\n\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1) if t)

    @staticmethod
    def _read_docx(data: bytes) -> str:
        from docx import Document as Docx
        d = Docx(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            parts += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        return "\n".join(parts)

    @staticmethod
    def _read_text(filename: str, data: bytes) -> str:
        if b"\x00" in data[:4096]:
            raise ValueError(f"Unsupported binary file: {filename}")
        for enc in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeError:
                continue
        raise ValueError(f"Could not decode file: {filename}")

    # ---------------- thread-scoped index ----------------
    def _store(self, thread_id: str) -> Chroma:
        name = ("att_" + re.sub(r"[^a-zA-Z0-9_-]", "_", thread_id))[:60].rstrip("_-") or "att_x"
        return Chroma(collection_name=name, embedding_function=self.embeddings,
                      persist_directory=settings.chroma_path)

    def _index_sync(self, thread_id: str, attachment_id: int, filename: str, text: str) -> None:
        pieces = self.splitter.split_text(text)
        docs = [Document(page_content=p, metadata={"attachment_id": attachment_id,
                                                   "filename": filename, "chunk_index": i})
                for i, p in enumerate(pieces) if p.strip()]
        store = self._store(thread_id)
        for i in range(0, len(docs), 4):  # small batches, same as repo ingestion
            store.add_documents(docs[i:i + 4], ids=[f"{attachment_id}:{d.metadata['chunk_index']}" for d in docs[i:i + 4]])

    async def index(self, thread_id: str, attachment_id: int, filename: str, text: str) -> None:
        await asyncio.to_thread(self._index_sync, thread_id, attachment_id, filename, text)

    # ---------------- prompt context ----------------
    async def build_context(self, thread_id: str, query: str, atts: List[Dict[str, Any]]) -> str:
        """Small thread attachments go in full; large ones go in as opening excerpt + query-relevant chunks."""
        if not atts:
            return ""
        names = ", ".join(f"{a['filename']} ({a['kind']})" for a in atts)
        header = f"Attachments in this conversation: {names}"
        total = sum(len(a["text"]) for a in atts)

        if total <= settings.attachment_inline_chars:
            blocks = [f"--- [ATTACHMENT: COMPLETE] {a['filename']} ---\n{a['text']}" for a in atts]
            return header + "\n\n" + "\n\n".join(blocks)

        blocks = [f"--- [ATTACHMENT: OPENING EXCERPT] {a['filename']} ---\n{a['text'][:OPENING_CHARS]}"
                  for a in atts[-5:]]
        try:
            docs = await asyncio.to_thread(
                lambda: self._store(thread_id).similarity_search(query, k=settings.attachment_top_k))
            docs = [d for d in docs if int(d.metadata.get("chunk_index", 0)) > 0]  # chunk 0 == opening excerpt
            docs.sort(key=lambda d: (d.metadata.get("filename", ""), int(d.metadata.get("chunk_index", 0))))
            blocks += [f"--- [ATTACHMENT: RELEVANT EXCERPT] {d.metadata.get('filename')} "
                       f"(chunk #{d.metadata.get('chunk_index')}) ---\n{d.page_content}" for d in docs]
        except Exception as exc:  # embeddings down: degrade to a larger head excerpt per file
            sys.stdout.write(f"[WARNING] Attachment retrieval failed: {exc}\n")
            per = max(OPENING_CHARS, settings.attachment_inline_chars // len(atts))
            blocks = [f"--- [ATTACHMENT: OPENING EXCERPT] {a['filename']} ---\n{a['text'][:per]}" for a in atts]
        return (header + "\nNOTE: attachments are large; only excerpts are shown, so never conclude something is absent.\n\n"
                + "\n\n".join(blocks))